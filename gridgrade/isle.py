"""Ham EIA dosyalarını site modeline çevirir: ABD ve eyalet aylık seriler, şirket metrikleri, güvenilirlik, ilçeler, notlar.
Sütunlar başlık metniyle bulunur; beklenen başlık yoksa işlem durur (canlı site eski hâliyle kalır)."""
import io
import json
import re
import unicodedata
import zipfile
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from . import ayar, notlar


class BicimHatasi(RuntimeError):
    pass


def slug(metin: str) -> str:
    s = unicodedata.normalize("NFKD", metin).encode("ascii", "ignore").decode().lower()
    s = s.replace("&", " and ")
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


def temiz_ad(ad: str) -> str:
    """EIA bazı adlara eyalet ekler: 'Borough of Chambersburg - (PA)' → 'Borough of Chambersburg'."""
    return re.sub(r"\s*-?\s*\([A-Z]{2}\)\s*$", "", re.sub(r"\s+", " ", ad)).strip()


def _sayi(x) -> float | None:
    if x is None or isinstance(x, str) and x.strip() in {"", ".", "NM", "W", "*"}:
        return None
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


def _satirlar(kaynak, sayfa: str) -> list[tuple]:
    wb = load_workbook(kaynak, read_only=True, data_only=True)
    if sayfa not in wb.sheetnames:
        raise BicimHatasi(f"'{sayfa}' sayfası yok: {wb.sheetnames}")
    return list(wb[sayfa].iter_rows(values_only=True))


def _baslik(satirlar: list[tuple], adet: int) -> list[str]:
    """Çok satırlı başlığı birleştirir: üst satırdaki grup adı sağa doğru taşınır."""
    genislik = max(len(r) for r in satirlar[:adet])
    katmanlar = []
    for r in satirlar[:adet]:
        dolu, son = [], ""
        for j in range(genislik):
            v = r[j] if j < len(r) else None
            if v not in (None, ""):
                son = str(v).strip()
            dolu.append(son)
        katmanlar.append(dolu)
    return [" | ".join(k[j] for k in katmanlar if k[j]).lower() for j in range(genislik)]


def _sutun(basliklar: list[str], *parcalar: str, atla: int = 0) -> int:
    adaylar = [j for j, b in enumerate(basliklar) if all(p in b for p in parcalar)]
    if len(adaylar) <= atla:
        raise BicimHatasi(f"başlık bulunamadı: {parcalar}")
    return adaylar[atla]


# ---- Aylık eyalet verisi (EIA-861M sales_revenue.xlsx) ----

SEKTORLER = {"res": "residential", "com": "commercial", "ind": "industrial", "all": "total"}


def aylik(yol: Path) -> dict:
    satirlar = _satirlar(yol, "Monthly-States")
    b = _baslik(satirlar, 3)
    yil, ay, ey, durum = _sutun(b, "year"), _sutun(b, "month"), _sutun(b, "state"), _sutun(b, "data status")
    s = {k: {m: _sutun(b, sek, m) for m in ("revenue", "sales", "customers")} for k, sek in SEKTORLER.items()}
    seri: dict[str, list] = defaultdict(list)
    for r in satirlar[3:]:
        if not r or r[ey] not in ayar.EYALETLER or not isinstance(r[yil], (int, float)):
            continue
        kayit = {"y": int(r[yil]), "m": int(r[ay]), "on": str(r[durum] or "").lower().startswith("prelim")}
        for k, c in s.items():
            kayit[k] = [_sayi(r[c["revenue"]]), _sayi(r[c["sales"]]), _sayi(r[c["customers"]])]  # bin $, MWh, adet
        seri[r[ey]].append(kayit)
    for v in seri.values():
        v.sort(key=lambda k: (k["y"], k["m"]))
    if len(seri) < 50:
        raise BicimHatasi(f"aylık veride yalnız {len(seri)} eyalet var")
    return dict(seri)


def fiyat(gelir_bin, satis_mwh) -> float | None:
    """¢/kWh = (bin $ × 1000 × 100) / (MWh × 1000)"""
    if not gelir_bin or not satis_mwh:
        return None
    return round(gelir_bin * 100 / satis_mwh, 2)


def fark_yuzde(yeni, eski) -> float | None:
    return round((yeni - eski) / eski * 100, 1) if yeni and eski else None


def abd_serisi(seri: dict) -> list[dict]:
    toplam: dict[tuple, dict] = {}
    for kayitlar in seri.values():
        for k in kayitlar:
            t = toplam.setdefault((k["y"], k["m"]), {"y": k["y"], "m": k["m"], "on": k["on"],
                                                     **{s: [0.0, 0.0, 0.0] for s in SEKTORLER}})
            for s in SEKTORLER:
                for i in range(3):
                    t[s][i] += k[s][i] or 0
    return [toplam[a] for a in sorted(toplam)]


def ozet(kayitlar: list[dict]) -> dict:
    """Son ay, bir yıl önceki aynı ay, son 12 ayın konut ortalama faturası ve tüketimi."""
    son = kayitlar[-1]
    once = next((k for k in kayitlar if k["y"] == son["y"] - 1 and k["m"] == son["m"]), None)
    on_iki = kayitlar[-12:]
    gelir = sum(k["res"][0] or 0 for k in on_iki)
    satis = sum(k["res"][1] or 0 for k in on_iki)
    abone = sum(k["res"][2] or 0 for k in on_iki) / len(on_iki)
    p = {s: fiyat(son[s][0], son[s][1]) for s in SEKTORLER}
    p_once = fiyat(once["res"][0], once["res"][1]) if once else None
    return {
        "y": son["y"], "m": son["m"], "on": son["on"], "fiyat": p,
        "yillik_degisim": round((p["res"] - p_once) / p_once * 100, 1) if p["res"] and p_once else None,
        "fatura": round(gelir * 1000 / abone / 12, 2) if abone else None,
        "tuketim": round(satis * 1000 / abone / 12) if abone else None,
        "fiyat_12ay": fiyat(gelir, satis),
    }


def yillik(kayitlar: list[dict]) -> dict[int, dict]:
    """Tam yılları (12 ay) toplar: konut fiyatı, aylık ortalama fatura ve tüketim."""
    yillar: dict[int, list] = defaultdict(list)
    for k in kayitlar:
        yillar[k["y"]].append(k)
    sonuc = {}
    for y, aylar in yillar.items():
        if len(aylar) < 12:
            continue
        gelir = sum(k["res"][0] or 0 for k in aylar)
        satis = sum(k["res"][1] or 0 for k in aylar)
        abone = sum(k["res"][2] or 0 for k in aylar) / 12
        sonuc[y] = {"fiyat": fiyat(gelir, satis), "fatura": round(gelir * 1000 / abone / 12, 2) if abone else None,
                    "tuketim": round(satis * 1000 / abone / 12) if abone else None, "abone": round(abone)}
    return sonuc


# ---- Eyalet elektrik üretimi kaynak payları (EIA-923 annual_generation_state) ----

KAYNAK_GRUBU = {"Natural Gas": "Natural gas", "Coal": "Coal", "Nuclear": "Nuclear", "Hydroelectric Conventional": "Hydro",
                "Wind": "Wind", "Solar Thermal and Photovoltaic": "Solar", "Petroleum": "Oil"}
KAYNAK_SIRASI = ["Natural gas", "Coal", "Nuclear", "Hydro", "Wind", "Solar", "Oil", "Other"]
KARSILASTIRMA_YILI = 10


def uretim(yol: Path) -> dict:
    """{eyalet: {yıl: {grup: pay%}}} — yalnız son yıl ve 10 yıl öncesi; pompaj depolama (net eksi) hariç."""
    wb = load_workbook(yol, read_only=True, data_only=True)
    satirlar = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    b = _baslik(satirlar[1:2], 1)
    yil, ey, ureten, kaynak, mwh = (_sutun(b, "year"), _sutun(b, "state"), _sutun(b, "type of producer"),
                                    _sutun(b, "energy source"), _sutun(b, "generation"))
    son = max(r[yil] for r in satirlar[2:] if isinstance(r[yil], int))
    hedef = {son, son - KARSILASTIRMA_YILI}
    toplam: dict = defaultdict(lambda: defaultdict(float))
    for r in satirlar[2:]:
        if (r[yil] not in hedef or r[ureten] != "Total Electric Power Industry" or r[kaynak] in ("Total", "Pumped Storage")
                or r[ey] not in ayar.EYALETLER):
            continue
        deger = _sayi(r[mwh]) or 0
        if deger > 0:
            toplam[(r[ey], r[yil])][KAYNAK_GRUBU.get(r[kaynak], "Other")] += deger
            toplam[("US", r[yil])][KAYNAK_GRUBU.get(r[kaynak], "Other")] += deger
    sonuc: dict = defaultdict(dict)
    for (e, y), gruplar in toplam.items():
        t = sum(gruplar.values())
        sonuc[e][str(y)] = {g: round(gruplar.get(g, 0) / t * 100, 1) for g in KAYNAK_SIRASI}
    return {"yil": son, "eyaletler": dict(sonuc)}


# ---- Yıllık şirket verisi (EIA-861 zip) ----

def _zip_sayfa(zip_yolu: Path, dosya_on_ek: str, sayfa: str) -> list[tuple]:
    z = zipfile.ZipFile(zip_yolu)
    ad = next((n for n in z.namelist() if n.startswith(dosya_on_ek) and n.endswith(".xlsx")), None)
    if not ad:
        raise BicimHatasi(f"{dosya_on_ek}*.xlsx zip içinde yok")
    return _satirlar(io.BytesIO(z.read(ad)), sayfa)


def sirketler(zip_yolu: Path) -> dict:
    satirlar = _zip_sayfa(zip_yolu, "Sales_Ult_Cust_", "States")
    b = _baslik(satirlar, 3)
    c = {"yil": _sutun(b, "data year"), "no": _sutun(b, "utility number"), "ad": _sutun(b, "utility name"),
         "parca": _sutun(b, "part"), "tur": _sutun(b, "service type"), "veri": _sutun(b, "data type"),
         "ey": _sutun(b, "state"), "sahip": _sutun(b, "ownership"),
         "gelir": _sutun(b, "residential", "revenues"), "satis": _sutun(b, "residential", "sales"),
         "abone": _sutun(b, "residential", "customers"),
         "t_gelir": _sutun(b, "total", "revenues"), "t_satis": _sutun(b, "total", "sales"),
         "t_abone": _sutun(b, "total", "customers")}
    # Bir şirket aynı eyalette birden çok dengeleme bölgesinde (BA) ayrı satır verir: satırlar toplanır
    toplam: dict[str, dict] = {}
    for r in satirlar[3:]:
        if not r or r[c["ey"]] not in ayar.EYALETLER or not isinstance(r[c["no"]], (int, float)):
            continue
        no = int(r[c["no"]])
        if no == 99999 or r[c["parca"]] != "A" or r[c["tur"]] != "Bundled" or r[c["sahip"]] in ("Behind the Meter", "", None):
            continue
        ey = r[c["ey"]]
        t = toplam.setdefault(f"{no}-{ey}", {"no": no, "ad": temiz_ad(str(r[c["ad"]])), "ey": ey, "yil": int(r[c["yil"]]),
                                            "sahip": str(r[c["sahip"]]).strip(), "tahmini": False,
                                            **{k: 0.0 for k in ("gelir", "satis", "abone", "t_gelir", "t_satis", "t_abone")}})
        t["tahmini"] |= r[c["veri"]] == "I"
        for k in ("gelir", "satis", "abone", "t_gelir", "t_satis", "t_abone"):
            t[k] += _sayi(r[c[k]]) or 0
    sonuc = {}
    for anahtar, t in toplam.items():
        if t["abone"] <= 0 or t["satis"] <= 0:
            continue
        sonuc[anahtar] = {
            "anahtar": anahtar, "no": t["no"], "ad": t["ad"], "ey": t["ey"], "yil": t["yil"], "sahip": t["sahip"],
            "tahmini": t["tahmini"], "abone": int(t["abone"]), "gelir": round(t["gelir"], 1), "satis": round(t["satis"]),
            "fiyat": fiyat(t["gelir"], t["satis"]),
            "fatura": round(t["gelir"] * 1000 / t["abone"] / 12, 2) if t["gelir"] else None,
            "tuketim": round(t["satis"] * 1000 / t["abone"] / 12),
            "toplam_abone": int(t["t_abone"]), "toplam_fiyat": fiyat(t["t_gelir"], t["t_satis"]),
        }
    if len(sonuc) < 1000:
        raise BicimHatasi(f"yalnız {len(sonuc)} şirket satırı")
    return sonuc


def guvenilirlik(zip_yolu: Path) -> dict:
    satirlar = _zip_sayfa(zip_yolu, "Reliability_", "Reliability_States")
    b = _baslik(satirlar, 3)
    no, ey = _sutun(b, "utility number"), _sutun(b, "state")
    ile = {k: _sutun(b, "ieee", "with major event days", k) for k in ("saidi", "saifi", "caidi")}
    haric = {k: _sutun(b, "ieee", "without major event days", k) for k in ("saidi", "saifi", "caidi")}
    sonuc = {}
    for r in satirlar[3:]:
        if not r or not isinstance(r[no], (int, float)) or r[ey] not in ayar.EYALETLER:
            continue
        d = {k: _sayi(r[j]) for k, j in haric.items()}
        d.update({f"{k}_ile": _sayi(r[j]) for k, j in ile.items()})
        if any(v is not None for v in d.values()):
            sonuc[f"{int(r[no])}-{r[ey]}"] = d
    return sonuc


def ozellikler(zip_yolu: Path) -> dict:
    """Akıllı sayaç oranı, konut için dinamik tarife programları, net metering (çatı güneşi) kurulumları."""
    sonuc: dict[str, dict] = defaultdict(lambda: {"ami": 0.0, "sayac": 0.0, "programlar": set(), "kayitli": 0.0,
                                                  "nm_kurulum": 0.0, "nm_mw": 0.0})

    def anahtar(r, no, ey):
        if r and isinstance(r[no], (int, float)) and r[ey] in ayar.EYALETLER:
            return f"{int(r[no])}-{r[ey]}"

    s = _zip_sayfa(zip_yolu, "Advanced_Meters_", "states")
    b = _baslik(s, 2)
    no, ey = _sutun(b, "utility number"), _sutun(b, "state")
    ami, top = _sutun(b, "ami", "residential"), _sutun(b, "total numbers of meters", "residential")
    for r in s[2:]:
        if k := anahtar(r, no, ey):
            sonuc[k]["ami"] += _sayi(r[ami]) or 0
            sonuc[k]["sayac"] += _sayi(r[top]) or 0

    s = _zip_sayfa(zip_yolu, "Dynamic_Pricing_", "Dynamic Pricing_States")
    b = _baslik(s, 3)
    no, ey = _sutun(b, "utility number"), _sutun(b, "state")
    kayit = _sutun(b, "customers enrolled", "residential")
    turler = {"Time-of-use": "time of use", "Real-time": "real time", "Variable peak": "variable peak",
              "Critical peak": "critical peak pricing", "Peak-time rebate": "critical peak rebate"}
    sutunlar = {ad: _sutun(b, parca, "residential") for ad, parca in turler.items()}
    for r in s[3:]:
        if k := anahtar(r, no, ey):
            sonuc[k]["kayitli"] += _sayi(r[kayit]) or 0
            sonuc[k]["programlar"] |= {ad for ad, j in sutunlar.items() if str(r[j] or "").strip().upper() == "Y"}

    s = _zip_sayfa(zip_yolu, "Net_Metering_", "States- Utility Level")
    b = _baslik(s, 3)
    no, ey = _sutun(b, "utility number"), _sutun(b, "state")
    kur, mw = _sutun(b, "photovoltaic", "installations", "residential"), _sutun(b, "photovoltaic", "capacity mw", "residential")
    for r in s[3:]:
        if k := anahtar(r, no, ey):
            sonuc[k]["nm_kurulum"] += _sayi(r[kur]) or 0
            sonuc[k]["nm_mw"] += _sayi(r[mw]) or 0

    return {k: {"ami_oran": round(v["ami"] / v["sayac"] * 100, 1) if v["sayac"] else None,
                "programlar": sorted(v["programlar"]), "kayitli": int(v["kayitli"]),
                "nm_kurulum": int(v["nm_kurulum"]), "nm_mw": round(v["nm_mw"], 1)} for k, v in sonuc.items()}


def ilceler(zip_yolu: Path) -> dict:
    satirlar = _zip_sayfa(zip_yolu, "Service_Territory_", "Counties_States")
    b = _baslik(satirlar, 1)
    no, ey, il = _sutun(b, "utility number"), _sutun(b, "state"), _sutun(b, "county")
    eslesme: dict[str, set] = defaultdict(set)
    for r in satirlar[1:]:
        if r and isinstance(r[no], (int, float)) and r[ey] in ayar.EYALETLER and r[il]:
            eslesme[f"{int(r[no])}-{r[ey]}"].add(str(r[il]).strip())
    return eslesme


# ---- Model ----

def model_kur(ham: Path = ayar.HAM) -> dict:
    seri = aylik(ham / "eia861m.xlsx")
    karisim = uretim(ham / "uretim.xlsx") if (ham / "uretim.xlsx").exists() else {"yil": None, "eyaletler": {}}
    abd = abd_serisi(seri)
    sirket = sirketler(ham / "eia861.zip")
    guv = guvenilirlik(ham / "eia861.zip")
    il_eslesme = ilceler(ham / "eia861.zip")
    ozel = ozellikler(ham / "eia861.zip")

    yil = next(iter(sirket.values()))["yil"]
    gecmis = {}
    for y in range(yil - ayar.GECMIS_YIL, yil):
        dosya = ham / f"eia861_{y}.zip"
        if dosya.exists():
            gecmis[y] = sirketler(dosya)
    for anahtar, s in sirket.items():
        s["gecmis"] = {str(y): {k: g[anahtar][k] for k in ("fiyat", "fatura", "abone")}
                       for y, g in sorted(gecmis.items()) if anahtar in g}
        onceki = s["gecmis"].get(str(yil - 1))
        s["degisim"] = fark_yuzde(s["fiyat"], onceki["fiyat"]) if onceki else None
        s["guv"] = guv.get(anahtar)
        s["ozellik"] = ozel.get(anahtar)
        s["ilceler"] = sorted(il_eslesme.get(anahtar, []))
    notlar.hesapla(sirket)

    # Aynı eyalette aynı ada sahip iki şirket olursa (nadir) numarayla ayrılır
    gorulen: dict[str, int] = defaultdict(int)
    for s in sorted(sirket.values(), key=lambda s: (s["ey"], -s["abone"])):
        k = f"{s['ey'].lower()}/{slug(s['ad'])}"
        gorulen[k] += 1
        s["yol"] = f"utilities/{k}" + (f"-{s['no']}" if gorulen[k] > 1 else "")

    eyalet = {}
    for ey, ad in ayar.EYALETLER.items():
        kayitlar = seri[ey]
        sirketleri = sorted((s for s in sirket.values() if s["ey"] == ey), key=lambda s: -s["abone"])
        eyalet[ey] = {"kod": ey, "ad": ad, "yol": f"states/{slug(ad)}", "ozet": ozet(kayitlar), "yillik": yillik(kayitlar),
                      "seri": [[k["y"], k["m"], fiyat(k["res"][0], k["res"][1]), fiyat(k["com"][0], k["com"][1]),
                                fiyat(k["ind"][0], k["ind"][1]), fiyat(k["all"][0], k["all"][1])] for k in kayitlar],
                      "sirketler": [s["anahtar"] for s in sirketleri]}
    for i, e in enumerate(sorted(eyalet.values(), key=lambda e: e["ozet"]["fiyat"]["res"] or 999), 1):
        e["sira"] = i  # 1 = en ucuz konut elektriği

    ilce = {}
    for s in sirket.values():
        for ad in s["ilceler"]:
            k = f"{s['ey'].lower()}/{slug(ad)}"
            d = ilce.setdefault(k, {"ad": ad, "ey": s["ey"], "yol": f"counties/{k}", "sirketler": []})
            d["sirketler"].append(s["anahtar"])
    for d in ilce.values():
        d["sirketler"].sort(key=lambda a: -sirket[a]["abone"])

    return {
        "abd": {"ozet": ozet(abd), "yillik": yillik(abd), "seri": [[k["y"], k["m"], fiyat(k["res"][0], k["res"][1]), fiyat(k["com"][0], k["com"][1]),
                                              fiyat(k["ind"][0], k["ind"][1]), fiyat(k["all"][0], k["all"][1])] for k in abd]},
        "eyaletler": eyalet, "sirketler": sirket, "ilceler": ilce, "uretim": karisim,
        "sirket_yili": next(iter(sirket.values()))["yil"],
        "kaynaklar": json.loads(ayar.KAYNAK_KAYDI.read_text(encoding="utf-8")) if ayar.KAYNAK_KAYDI.exists() else {},
    }


def kaydet(model: dict, yol: Path = ayar.MODEL) -> None:
    yol.parent.mkdir(parents=True, exist_ok=True)
    yol.write_text(json.dumps(model, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
